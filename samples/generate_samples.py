import os
import docx
import pymupdf as fitz

SAMPLE_QUESTIONS = """1. Which of the following are operating systems?
A. Windows
B. Linux
C. Chrome
D. Ubuntu
Answer: A, B, D
Explanation: Windows, Linux, and Ubuntu are operating systems. Google Chrome is a web browser.

2. Which of the following are relational database management systems?
A. PostgreSQL
B. MySQL
C. Redis
D. Oracle Database
Answer: A, B, D
Explanation: PostgreSQL, MySQL, and Oracle are relational (SQL) databases. Redis is an in-memory key-value store.

3. Which of the following are application layer protocols in the OSI / TCP-IP model?
A. HTTP
B. FTP
C. TCP
D. DNS
Answer: A, B, D
Explanation: HTTP, FTP, and DNS operate at the Application Layer (Layer 7). TCP operates at the Transport Layer (Layer 4).

4. Which of the following are network security devices or systems?
A. Firewall
B. IDS (Intrusion Detection System)
C. IPS (Intrusion Prevention System)
D. Compiler
Answer: A, B, C
Explanation: Firewalls, IDS, and IPS are security systems designed to protect networks. A compiler translates programming language source code.

5. Which of the following data structures have O(1) average time complexity for insertion?
A. Hash Table
B. Singly Linked List (at head)
C. Binary Search Tree (unbalanced worst case)
D. Array List (insertion at arbitrary middle index)
Answer: A, B
Explanation: Hash tables average O(1) insertions, and inserting at the head of a linked list is O(1). BST insertions take O(log n) average and O(n) worst-case.

6. Which of the following are symmetric encryption algorithms?
A. AES (Advanced Encryption Standard)
B. DES (Data Encryption Standard)
C. RSA
D. Blowfish
Answer: A, B, D
Explanation: AES, DES, and Blowfish use symmetric keys. RSA is an asymmetric (public-key) cryptosystem.

7. Which of the following are recognized cloud computing service models?
A. IaaS (Infrastructure as a Service)
B. PaaS (Platform as a Service)
C. SaaS (Software as a Service)
D. CaaS (Compiler as a Service)
Answer: A, B, C
Explanation: The NIST cloud model defines IaaS, PaaS, and SaaS as the primary service delivery models.
"""


def create_txt():
    txt_path = os.path.join(os.path.dirname(__file__), "sample_exam.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(SAMPLE_QUESTIONS.strip())
    print(f"Created {txt_path}")


def create_docx():
    docx_path = os.path.join(os.path.dirname(__file__), "sample_exam.docx")
    doc = docx.Document()
    doc.add_heading("Computer Science & Networking MSQ Practice Exam", level=1)
    doc.add_paragraph("Comprehensive Multiple Select Questions study material.")

    for line in SAMPLE_QUESTIONS.strip().split("\n"):
        line_str = line.strip()
        if not line_str:
            doc.add_paragraph("")
        elif line_str.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.")):
            p = doc.add_paragraph()
            run = p.add_run(line_str)
            run.bold = True
        elif line_str.startswith("Answer:"):
            p = doc.add_paragraph()
            run = p.add_run(line_str)
            run.italic = True
        elif line_str.startswith("Explanation:"):
            p = doc.add_paragraph()
            p.add_run(line_str)
        else:
            doc.add_paragraph(line_str)

    doc.save(docx_path)
    print(f"Created {docx_path}")


def create_pdf():
    pdf_path = os.path.join(os.path.dirname(__file__), "sample_exam.pdf")
    doc = fitz.open()

    # 595 x 842 is standard A4 points
    page = doc.new_page(width=595, height=842)
    y = 50
    page.insert_text((50, y), "Computer Science & Networking MSQ Practice Exam", fontsize=14, fontname="helv", color=(0, 0.2, 0.6))
    y += 20
    page.insert_text((50, y), "Official Study Material & Exam Questions", fontsize=11, fontname="helv", color=(0.3, 0.3, 0.3))
    y += 25

    lines = SAMPLE_QUESTIONS.strip().split("\n")
    for line in lines:
        line_str = line.strip()
        if y > 790:
            page = doc.new_page(width=595, height=842)
            y = 50

        if not line_str:
            y += 8
            continue

        if line_str.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.")):
            y += 6
            page.insert_text((50, y), line_str, fontsize=10, fontname="hebo", color=(0.1, 0.1, 0.1))
        elif line_str.startswith("Answer:"):
            page.insert_text((50, y), line_str, fontsize=9.5, fontname="heit", color=(0, 0.5, 0.1))
        elif line_str.startswith("Explanation:"):
            page.insert_text((50, y), line_str, fontsize=9, fontname="helv", color=(0.4, 0.4, 0.4))
        else:
            page.insert_text((60, y), line_str, fontsize=9.5, fontname="helv", color=(0.2, 0.2, 0.2))

        y += 14

    doc.save(pdf_path)
    doc.close()
    print(f"Created {pdf_path}")


if __name__ == "__main__":
    create_txt()
    create_docx()
    create_pdf()
